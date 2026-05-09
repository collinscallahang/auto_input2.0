from __future__ import annotations

import argparse
import sys
import time
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from huolala_quote_tool.browser import HuolalaClient
from huolala_quote_tool.excel_model import (
    clear_address_mismatch,
    create_output_path,
    detect_workbook,
    iter_data_rows,
    load_vehicle_rules,
    mark_address_mismatch,
    mark_failure,
    numeric_or_original,
    value_at,
    writable_cell,
    write_success,
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run a real Huolala batch test on sample workbook rows.")
    parser.add_argument("source", type=Path, help="Sample workbook to test.")
    parser.add_argument("--rows", type=int, default=10, help="Number of data rows to process.")
    parser.add_argument("--vehicles", type=int, default=3, help="Number of detected vehicle columns per row.")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path.cwd() / "real_batch_output",
        help="Directory for result workbook, log, and screenshots.",
    )
    parser.add_argument(
        "--keep-open-seconds",
        type=int,
        default=0,
        help="Keep the browser open after the run so the tester can inspect the page.",
    )
    return parser.parse_args()


def failure_message(step: str, exc: Exception) -> str:
    raw = " ".join(str(exc).split())
    analysis = "原因待人工确认"
    if "地址识别" in raw or "intercepts pointer events" in raw or "选择准确地址" in raw:
        analysis = "地址识别候选没有匹配 Excel 中的供应商/发货地址/到货地址，或弹窗没有正常关闭"
    elif "无法选择车长" in raw:
        analysis = "网页没有成功切到目标车长，可能页面未加载完成、登录失效或车长文案变化"
    elif "车型要求未切换" in raw or "无法选择车型要求" in raw:
        analysis = "车型要求没有切换到规则指定状态，可能有错误勾选残留或网页限制至少保留一项"
    elif "运费一口价" in raw or "计价中" in raw:
        analysis = "网页没有生成可读取的运费一口价，可能计价仍在加载或当前路线暂不支持报价"
    elif "总里程" in raw:
        analysis = "网页没有生成可读取的总里程，可能地址未被准确识别"
    elif "Timeout" in raw:
        analysis = "网页响应超时，可能网络慢、登录过期或页面控件被弹窗挡住"
    return f"{step}失败：{analysis}。原始提示：{raw[:260]}"


def compare_and_mark_addresses(
    ws,
    row_idx: int,
    supplier: str,
    origin_col,
    destination_col,
    expected_origin: str,
    expected_destination: str,
    actual_origin: str,
    actual_destination: str,
    log,
) -> list[dict[str, str | int]]:
    mismatches: list[dict[str, str | int]] = []
    checks = [
        ("发货地址", origin_col, expected_origin, actual_origin),
        ("到货地址", destination_col, expected_destination, actual_destination),
    ]
    for address_type, column, expected, actual in checks:
        expected_text = (expected or "").strip()
        actual_text = (actual or "").strip()
        cell = writable_cell(ws, row_idx, column.index)
        if expected_text == actual_text:
            clear_address_mismatch(cell)
            log(f"row {row_idx} {address_type} exact_match=true")
            continue

        expected_display = expected_text or "空"
        actual_display = actual_text or "未读取到"
        message = f"{address_type}不一致；Excel={expected_display}；网页参与计价={actual_display}"
        mark_address_mismatch(cell, message)
        mismatches.append(
            {
                "row": row_idx,
                "supplier": supplier or "",
                "address_type": address_type,
                "excel_address": expected_display,
                "actual_address": actual_display,
            }
        )
        log(
            f"row {row_idx} address_mismatch {address_type}: "
            f"excel={expected_display}; actual={actual_display}; needs_manual_fix=true"
        )
    return mismatches


def write_address_mismatch_summary_sheet(wb, mismatches: list[dict[str, str | int]]) -> None:
    sheet_name = "地址不一致汇总"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    if not mismatches:
        return

    ws = wb.create_sheet(sheet_name)
    ws.append(["行号", "供应商", "地址类型", "Excel地址", "网页参与计价地址", "处理建议"])
    for item in mismatches:
        ws.append(
            [
                item["row"],
                item["supplier"],
                item["address_type"],
                item["excel_address"],
                item["actual_address"],
                "请人工核对并修正 Excel 地址后重跑",
            ]
        )
    for column, width in {
        "A": 10,
        "B": 28,
        "C": 14,
        "D": 46,
        "E": 46,
        "F": 30,
    }.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"


def main() -> None:
    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    log_path = args.output_dir / f"batch_{time.strftime('%Y%m%d_%H%M%S')}.log"

    def log(message: str) -> None:
        line = f"[{time.strftime('%H:%M:%S')}] {message}"
        print(line, flush=True)
        with log_path.open("a", encoding="utf-8") as f:
            f.write(line + "\n")

    rules = load_vehicle_rules()
    detection = detect_workbook(args.source, rules=rules)
    if not detection.can_start:
        raise SystemExit(f"字段识别不完整：{'；'.join(detection.warnings)}")

    wb = load_workbook(args.source, data_only=False)
    ws = wb[detection.sheet_name]
    rows = iter_data_rows(ws, detection)[: args.rows]
    vehicle_columns = detection.vehicle_columns[: args.vehicles]
    output_path = args.output_dir / create_output_path(args.source).name

    origin_col = detection.required_columns["origin"]
    destination_col = detection.required_columns["destination"]
    distance_col = detection.required_columns["distance"]
    supplier_col = detection.required_columns["supplier"]
    assert origin_col and destination_col and distance_col and supplier_col

    # Start from blank output cells so the result proves this run wrote them.
    for row_idx in rows:
        writable_cell(ws, row_idx, distance_col.index).value = None
        writable_cell(ws, row_idx, distance_col.index).comment = None
        for vehicle_col in vehicle_columns:
            cell = writable_cell(ws, row_idx, vehicle_col.index)
            cell.value = None
            cell.comment = None
    wb.save(output_path)

    log(f"source={args.source}")
    log(f"output={output_path}")
    log(f"rows={len(rows)} vehicles={len(vehicle_columns)}")

    client = HuolalaClient(log=lambda message: log(f"browser: {message}"))
    address_mismatches: list[dict[str, str | int]] = []
    previous_origin = ""
    previous_destination = ""
    try:
        client.start()
        client.open_order_page()

        for index, row_idx in enumerate(rows, start=1):
            supplier = str(value_at(ws, row_idx, supplier_col.index) or "").strip()
            origin = str(value_at(ws, row_idx, origin_col.index) or "").strip()
            destination = str(value_at(ws, row_idx, destination_col.index) or "").strip()
            log(f"row {row_idx} ({index}/{len(rows)}) start")
            log(f"row {row_idx} supplier={supplier}")
            log(f"row {row_idx} origin={origin}")
            log(f"row {row_idx} destination={destination}")

            if not origin or not destination:
                message = "发货地址或到货地址为空"
                mark_failure(writable_cell(ws, row_idx, distance_col.index), message)
                for vehicle_col in vehicle_columns:
                    mark_failure(writable_cell(ws, row_idx, vehicle_col.index), message)
                wb.save(output_path)
                log(f"row {row_idx} skipped: {message}")
                continue

            try:
                address_result = client.fill_addresses(
                    origin,
                    destination,
                    reuse_if_same=(origin == previous_origin and destination == previous_destination),
                    supplier=supplier,
                )
                previous_origin = origin
                previous_destination = destination
                actual_origin = (address_result.actual_origin or "").strip()
                actual_destination = (address_result.actual_destination or "").strip()
                log(f"row {row_idx} actual_origin={actual_origin or '未读取到'}")
                log(f"row {row_idx} actual_destination={actual_destination or '未读取到'}")
                address_mismatches.extend(
                    compare_and_mark_addresses(
                        ws,
                        row_idx,
                        supplier,
                        origin_col,
                        destination_col,
                        origin,
                        destination,
                        actual_origin,
                        actual_destination,
                        log,
                    )
                )
            except Exception as exc:
                message = failure_message("地址填写", exc)
                mark_failure(writable_cell(ws, row_idx, distance_col.index), message)
                for vehicle_col in vehicle_columns:
                    mark_failure(writable_cell(ws, row_idx, vehicle_col.index), message)
                wb.save(output_path)
                log(f"row {row_idx} failed={message}")
                continue

            for vehicle_col in vehicle_columns:
                cell = writable_cell(ws, row_idx, vehicle_col.index)
                try:
                    assert vehicle_col.rule is not None
                    price = client.quote_vehicle(vehicle_col.rule)
                    write_success(cell, numeric_or_original(price))
                    log(f"row {row_idx} {vehicle_col.header} price={price:g}")
                except Exception as exc:
                    message = failure_message(f"{vehicle_col.header} 报价", exc)
                    mark_failure(cell, message)
                    log(f"row {row_idx} {vehicle_col.header} failed={message}")

            try:
                distance = client.read_distance()
                write_success(writable_cell(ws, row_idx, distance_col.index), numeric_or_original(distance))
                log(f"row {row_idx} distance={distance:g}")
            except Exception as exc:
                message = failure_message("读取总里程", exc)
                mark_failure(writable_cell(ws, row_idx, distance_col.index), message)
                log(f"row {row_idx} distance failed={message}")

            wb.save(output_path)
            log(f"row {row_idx} saved")

        page = client.page
        if page is not None:
            screenshot_path = args.output_dir / "batch_last_page.png"
            text_path = args.output_dir / "batch_last_page_body.txt"
            page.screenshot(path=str(screenshot_path), full_page=True)
            text_path.write_text(page.locator("body").inner_text(timeout=5000), encoding="utf-8")
            log(f"screenshot={screenshot_path}")
            log(f"page_text={text_path}")

        write_address_mismatch_summary_sheet(wb, address_mismatches)
        wb.save(output_path)
        if address_mismatches:
            log(f"address_mismatch_summary count={len(address_mismatches)} needs_manual_fix=true")
            for item in address_mismatches:
                log(
                    f"address_mismatch_summary row={item['row']} supplier={item['supplier']} "
                    f"type={item['address_type']} excel={item['excel_address']} actual={item['actual_address']}"
                )
        else:
            log("address_mismatch_summary count=0")
        log("status=ok")
    finally:
        if args.keep_open_seconds > 0:
            log(f"keeping_browser_open_seconds={args.keep_open_seconds}")
            time.sleep(args.keep_open_seconds)
        client.stop()


if __name__ == "__main__":
    main()

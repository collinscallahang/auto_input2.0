from __future__ import annotations

import argparse
import sys
import time
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from huolala_quote_tool.browser import HuolalaClient
from huolala_quote_tool.excel_model import VehicleRule
from huolala_quote_tool.parsers import parse_fixed_price


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Diagnose Huolala 9.6m quote differences.")
    parser.add_argument("source", type=Path)
    parser.add_argument("--row", type=int, default=3)
    parser.add_argument("--keep-open-seconds", type=int, default=1800)
    parser.add_argument("--output-dir", type=Path, default=Path.cwd() / "real_batch_output" / "diagnostics")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    log_path = args.output_dir / f"diagnose_96_{time.strftime('%Y%m%d_%H%M%S')}.log"

    def log(message: str) -> None:
        line = f"[{time.strftime('%H:%M:%S')}] {message}"
        print(line, flush=True)
        with log_path.open("a", encoding="utf-8") as f:
            f.write(line + "\n")

    wb = load_workbook(args.source, data_only=True)
    ws = wb.active
    row = args.row
    origin = str(ws.cell(row, 3).value).strip()
    destination = str(ws.cell(row, 8).value).strip()
    expected = ws.cell(row, 5).value
    log(f"source={args.source}")
    log(f"row={row} expected_9.6={expected}")
    log(f"origin={origin}")
    log(f"destination={destination}")

    client = HuolalaClient(log=lambda message: log(f"browser: {message}"), timeout_ms=30000)
    client.start()
    client.open_order_page()
    client.fill_addresses(origin, destination, reuse_if_same=False)

    def read_price_section() -> str:
        body = client.page.locator("body").inner_text(timeout=3000)
        start = max(0, body.find("用车时间"))
        end = body.find("下一步")
        if end == -1:
            end = len(body)
        return body[start:end].strip()

    box = "厢式货车"
    flying = "飞翼车"
    scenarios = [
        ("9米6 only_box", "9米6", (box,)),
        ("9米6 box_and_flying", "9米6", (box, flying)),
        ("13米 box_and_flying", "13米", (box, flying)),
        ("17米5 box_and_flying", "17米5", (box, flying)),
    ]
    for name, length, requirements in scenarios:
        try:
            rule = VehicleRule(name=name, excel_headers=(), car_length=length, vehicle_requirements=requirements)
            price = client.quote_vehicle(rule)
            checked = client._checked_vehicle_requirements()
            log(f"scenario={name} requirements={list(requirements)} price={price:g} checked={checked}")
        except Exception as exc:
            log(f"scenario={name} failed={exc}")

    screenshot = args.output_dir / f"diagnose_96_last_page_{time.strftime('%Y%m%d_%H%M%S')}.png"
    text_path = args.output_dir / f"diagnose_96_last_page_{time.strftime('%Y%m%d_%H%M%S')}.txt"
    client.page.screenshot(path=str(screenshot), full_page=True)
    text_path.write_text(read_price_section(), encoding="utf-8")
    log(f"screenshot={screenshot}")
    log(f"price_section={text_path}")
    log(f"leaving_browser_open_seconds={args.keep_open_seconds}")
    time.sleep(args.keep_open_seconds)


if __name__ == "__main__":
    main()

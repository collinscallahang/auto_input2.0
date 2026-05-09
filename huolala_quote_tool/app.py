from __future__ import annotations

import os
import queue
import subprocess
import sys
import threading
import time
import traceback
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, scrolledtext, ttk

from openpyxl import load_workbook

from .browser import HuolalaClient
from .excel_model import (
    DetectionResult,
    VehicleRule,
    create_output_path,
    detect_workbook,
    iter_data_rows,
    load_vehicle_rules,
    clear_address_mismatch,
    mark_address_mismatch,
    mark_failure,
    numeric_or_original,
    should_process_cell,
    value_at,
    writable_cell,
    write_success,
)
from .paths import logs_dir


RERUN_STRATEGIES = {
    "只补空白价格": "empty",
    "全部重跑": "all",
    "只重跑失败项": "failed",
}


class HuolalaQuoteApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("货拉拉 Excel 批量报价工具")
        self.root.geometry("1120x760")
        self.root.minsize(960, 640)

        self.selected_path: Path | None = None
        self.detection: DetectionResult | None = None
        self.worker: threading.Thread | None = None
        self.stop_event = threading.Event()
        self.pause_event = threading.Event()
        self.pause_event.set()
        self.events: queue.Queue[tuple] = queue.Queue()
        self.log_path = logs_dir() / f"huolala_quote_tool_{time.strftime('%Y%m%d_%H%M%S')}.log"
        self.log_path.write_text("", encoding="utf-8")
        self.last_output_path: Path | None = None

        self._build_ui()
        self._set_running(False)
        self._log(f"日志文件：{self.log_path}")
        self.root.after(300, self._bring_to_front)
        self.root.after(100, self._drain_events)
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)

    def _build_ui(self) -> None:
        self.root.columnconfigure(0, weight=1)
        self.root.rowconfigure(1, weight=1)
        self.root.rowconfigure(4, weight=1)

        file_frame = ttk.Frame(self.root, padding=(12, 12, 12, 6))
        file_frame.grid(row=0, column=0, sticky="ew")
        file_frame.columnconfigure(1, weight=1)

        ttk.Button(file_frame, text="选择 Excel", command=self._choose_file).grid(row=0, column=0, padx=(0, 8))
        self.file_var = tk.StringVar()
        ttk.Entry(file_frame, textvariable=self.file_var, state="readonly").grid(row=0, column=1, sticky="ew")
        self.sheet_var = tk.StringVar(value="工作表：-")
        ttk.Label(file_frame, textvariable=self.sheet_var).grid(row=0, column=2, padx=(10, 0))

        preview_frame = ttk.LabelFrame(self.root, text="字段识别预览", padding=(12, 8))
        preview_frame.grid(row=1, column=0, sticky="nsew", padx=12, pady=(0, 8))
        preview_frame.columnconfigure(0, weight=1)
        preview_frame.rowconfigure(0, weight=1)

        columns = ("type", "column", "header", "mapping", "status")
        self.preview_tree = ttk.Treeview(preview_frame, columns=columns, show="headings", height=12)
        self.preview_tree.heading("type", text="类型")
        self.preview_tree.heading("column", text="Excel列")
        self.preview_tree.heading("header", text="表头")
        self.preview_tree.heading("mapping", text="识别结果/规则")
        self.preview_tree.heading("status", text="状态")
        self.preview_tree.column("type", width=120, anchor="w")
        self.preview_tree.column("column", width=80, anchor="center")
        self.preview_tree.column("header", width=260, anchor="w")
        self.preview_tree.column("mapping", width=380, anchor="w")
        self.preview_tree.column("status", width=140, anchor="w")
        self.preview_tree.grid(row=0, column=0, sticky="nsew")

        preview_scroll = ttk.Scrollbar(preview_frame, orient="vertical", command=self.preview_tree.yview)
        preview_scroll.grid(row=0, column=1, sticky="ns")
        self.preview_tree.configure(yscrollcommand=preview_scroll.set)

        control_frame = ttk.Frame(self.root, padding=(12, 0, 12, 8))
        control_frame.grid(row=2, column=0, sticky="ew")
        control_frame.columnconfigure(5, weight=1)

        ttk.Label(control_frame, text="复跑策略").grid(row=0, column=0, padx=(0, 8))
        self.strategy_var = tk.StringVar(value="只补空白价格")
        self.strategy_combo = ttk.Combobox(
            control_frame,
            textvariable=self.strategy_var,
            values=list(RERUN_STRATEGIES.keys()),
            state="readonly",
            width=16,
        )
        self.strategy_combo.grid(row=0, column=1, padx=(0, 12))

        self.start_button = ttk.Button(control_frame, text="开始", command=self._start)
        self.start_button.grid(row=0, column=2, padx=(0, 8))
        self.pause_button = ttk.Button(control_frame, text="暂停", command=self._toggle_pause)
        self.pause_button.grid(row=0, column=3, padx=(0, 8))
        self.stop_button = ttk.Button(control_frame, text="停止", command=self._stop)
        self.stop_button.grid(row=0, column=4, padx=(0, 12))

        self.progress_var = tk.DoubleVar(value=0)
        self.progress = ttk.Progressbar(control_frame, variable=self.progress_var, maximum=100)
        self.progress.grid(row=0, column=5, sticky="ew", padx=(0, 8))
        self.progress_label_var = tk.StringVar(value="未开始")
        ttk.Label(control_frame, textvariable=self.progress_label_var, width=22).grid(row=0, column=6, sticky="e")

        result_frame = ttk.LabelFrame(self.root, text="结果文件", padding=(12, 8))
        result_frame.grid(row=3, column=0, sticky="ew", padx=12, pady=(0, 8))
        result_frame.columnconfigure(1, weight=1)

        ttk.Label(result_frame, text="当前结果").grid(row=0, column=0, padx=(0, 8))
        self.result_file_var = tk.StringVar(value="任务完成后会显示结果文件")
        ttk.Entry(result_frame, textvariable=self.result_file_var, state="readonly").grid(row=0, column=1, sticky="ew")
        self.open_result_button = ttk.Button(result_frame, text="打开结果文件", command=self._open_result_file)
        self.open_result_button.grid(row=0, column=2, padx=(8, 0))
        self.open_result_folder_button = ttk.Button(result_frame, text="打开所在文件夹", command=self._open_result_folder)
        self.open_result_folder_button.grid(row=0, column=3, padx=(8, 0))

        log_frame = ttk.LabelFrame(self.root, text="进度和日志", padding=(12, 8))
        log_frame.grid(row=4, column=0, sticky="nsew", padx=12, pady=(0, 12))
        log_frame.columnconfigure(0, weight=1)
        log_frame.rowconfigure(0, weight=1)

        self.log_text = scrolledtext.ScrolledText(log_frame, height=12, wrap="word", state="disabled")
        self.log_text.grid(row=0, column=0, sticky="nsew")

    def _choose_file(self) -> None:
        path = filedialog.askopenfilename(
            title="选择 Excel",
            filetypes=[("Excel 文件", "*.xlsx *.xlsm"), ("所有文件", "*.*")],
        )
        if not path:
            return
        self.selected_path = Path(path)
        self.file_var.set(str(self.selected_path))
        self.last_output_path = None
        self.result_file_var.set("任务完成后会显示结果文件")
        self._refresh_result_buttons(running=False)
        self._log(f"开始识别：{self.selected_path}")
        try:
            rules = load_vehicle_rules()
            self.detection = detect_workbook(self.selected_path, rules=rules)
        except Exception as exc:
            self.detection = None
            self.sheet_var.set("工作表：-")
            self._clear_preview()
            self._set_running(False)
            messagebox.showerror("识别失败", str(exc), parent=self.root)
            self._log(traceback.format_exc())
            return
        self._show_detection(self.detection)
        self._set_running(False)

    def _clear_preview(self) -> None:
        for item in self.preview_tree.get_children():
            self.preview_tree.delete(item)

    def _show_detection(self, detection: DetectionResult) -> None:
        self._clear_preview()
        self.sheet_var.set(f"工作表：{detection.sheet_name}；表头行：{detection.header_row}")

        for role, column in detection.required_columns.items():
            if column:
                self.preview_tree.insert(
                    "",
                    "end",
                    values=(column.label, column.letter, column.header, f"置信度 {column.confidence:.0%}", "已识别"),
                )
            else:
                self.preview_tree.insert("", "end", values=(role, "-", "-", "未识别", "缺失"))

        for column in detection.vehicle_columns:
            self.preview_tree.insert(
                "",
                "end",
                values=(
                    "车型价格列",
                    column.letter,
                    column.header,
                    column.message,
                    "会处理",
                ),
            )

        for column in detection.unmatched_vehicle_headers:
            self.preview_tree.insert(
                "",
                "end",
                values=("疑似车型列", column.letter, column.header, column.message, "跳过"),
            )

        if detection.warnings:
            self._log("识别提示：")
            for warning in detection.warnings:
                self._log(f"  - {warning}")
        else:
            self._log("识别完成，必需字段和车型规则均已匹配。")

    def _start(self) -> None:
        if not self.detection:
            messagebox.showwarning("未选择 Excel", "请先选择 Excel 并确认字段识别结果。", parent=self.root)
            return
        if not self.detection.can_start:
            missing = "、".join(self.detection.missing_required_roles)
            reason = f"缺少必需字段：{missing}" if missing else "没有匹配车型规则的价格列"
            messagebox.showwarning("不能开始", reason, parent=self.root)
            return
        if self.worker and self.worker.is_alive():
            return

        self.stop_event.clear()
        self.pause_event.set()
        self.pause_button.configure(text="暂停")
        self.last_output_path = None
        self.result_file_var.set("正在生成结果文件...")
        self._refresh_result_buttons(running=True)
        strategy = RERUN_STRATEGIES[self.strategy_var.get()]
        self._set_running(True)
        self.worker = threading.Thread(target=self._run_job, args=(self.detection, strategy), daemon=True)
        self.worker.start()

    def _toggle_pause(self) -> None:
        if not self.worker or not self.worker.is_alive():
            return
        if self.pause_event.is_set():
            self.pause_event.clear()
            self.pause_button.configure(text="继续")
            self._log("已暂停，当前行完成后会停住等待。")
        else:
            self.pause_event.set()
            self.pause_button.configure(text="暂停")
            self._log("继续执行。")

    def _stop(self) -> None:
        if not self.worker or not self.worker.is_alive():
            return
        self.stop_event.set()
        self.pause_event.set()
        self._log("已请求停止，程序会在安全位置退出。")

    def _run_job(self, detection: DetectionResult, strategy: str) -> None:
        client: HuolalaClient | None = None
        output_path: Path | None = None
        try:
            wb = load_workbook(detection.workbook_path)
            ws = wb[detection.sheet_name]
            output_path = create_output_path(detection.workbook_path)
            wb.save(output_path)
            self._event("result_path", str(output_path))
            rows = iter_data_rows(ws, detection)
            if not rows:
                self._event("log", "没有找到可处理的数据行。")
                return

            self._event("log", f"任务启动：共 {len(rows)} 行，结果会写入副本，不覆盖原文件。")
            self._event("log", f"结果文件：{output_path}")
            self._event("log", f"复跑策略：{strategy}")
            self._event("log", "车型规则：")
            for column in detection.vehicle_columns:
                assert column.rule is not None
                requirements = "、".join(column.rule.vehicle_requirements) or "未指定"
                self._event(
                    "log",
                    f"  {column.letter}:{column.header} -> 车长 {column.rule.car_length}，车型要求 {requirements}",
                )
            self._event("progress", 0, len(rows), "准备浏览器")
            client = HuolalaClient(log=lambda message: self._event("log", message))
            client.start()
            client.open_order_page()

            ok = self._ask_user(
                "登录确认",
                "浏览器已经打开货拉拉同城下单页。\n\n请在浏览器里完成登录，并确认页面停留在同城下单页后点击“确定”继续。",
            )
            if not ok:
                self._event("log", "用户取消执行。")
                return

            address_mismatches: list[dict[str, str | int]] = []
            previous_origin = ""
            previous_destination = ""
            for index, row_idx in enumerate(rows, start=1):
                if self.stop_event.is_set():
                    self._event("log", "任务已停止。")
                    break
                self._wait_if_paused()

                origin_col = detection.required_columns["origin"]
                destination_col = detection.required_columns["destination"]
                distance_col = detection.required_columns["distance"]
                supplier_col = detection.required_columns["supplier"]
                assert origin_col and destination_col and distance_col and supplier_col

                supplier = str(value_at(ws, row_idx, supplier_col.index) or "").strip()
                origin = str(value_at(ws, row_idx, origin_col.index) or "").strip()
                destination = str(value_at(ws, row_idx, destination_col.index) or "").strip()
                self._event("log", f"第 {row_idx} 行：开始处理")
                self._event("log", f"第 {row_idx} 行：供应商={supplier or '空'}")
                self._event("log", f"第 {row_idx} 行：发货地址={origin or '空'}")
                self._event("log", f"第 {row_idx} 行：到货地址={destination or '空'}")

                distance_cell = writable_cell(ws, row_idx, distance_col.index)
                vehicle_cells = [
                    (column, writable_cell(ws, row_idx, column.index))
                    for column in detection.vehicle_columns
                    if should_process_cell(writable_cell(ws, row_idx, column.index), strategy)
                ]
                process_distance = should_process_cell(distance_cell, strategy)

                if not origin or not destination:
                    message = "发货地址或到货地址为空"
                    if process_distance:
                        mark_failure(distance_cell, message)
                    for _, cell in vehicle_cells:
                        mark_failure(cell, message)
                    wb.save(output_path)
                    self._event("log", f"第 {row_idx} 行跳过：{message}")
                    self._event("progress", index, len(rows), f"{index}/{len(rows)}")
                    continue

                if not process_distance and not vehicle_cells:
                    self._event("log", f"第 {row_idx} 行无须处理，按当前复跑策略跳过。")
                    wb.save(output_path)
                    self._event("progress", index, len(rows), f"{index}/{len(rows)}")
                    continue

                try:
                    self._event("log", f"第 {row_idx} 行：正在填写地址到货拉拉页面")
                    address_result = client.fill_addresses(
                        origin,
                        destination,
                        reuse_if_same=(origin == previous_origin and destination == previous_destination),
                        supplier=supplier,
                    )
                    previous_origin = origin
                    previous_destination = destination
                    actual_origin = (address_result.actual_origin or "").strip()
                    actual_destination = (address_result.actual_destination or "").strip()
                    self._event("log", f"第 {row_idx} 行：网页最终参与计价发货地址={actual_origin or '未读取到'}")
                    self._event("log", f"第 {row_idx} 行：网页最终参与计价到货地址={actual_destination or '未读取到'}")
                    address_mismatches.extend(
                        self._compare_and_mark_addresses(
                            ws,
                            row_idx,
                            supplier,
                            origin_col,
                            destination_col,
                            origin,
                            destination,
                            actual_origin,
                            actual_destination,
                        )
                    )
                    self._event("log", f"第 {row_idx} 行：地址填写完成")
                except Exception as exc:
                    message = self._failure_message("地址填写", exc)
                    if process_distance:
                        mark_failure(distance_cell, message)
                    for _, cell in vehicle_cells:
                        mark_failure(cell, message)
                    wb.save(output_path)
                    self._event("log", f"第 {row_idx} 行失败：{message}")
                    self._event("progress", index, len(rows), f"{index}/{len(rows)}")
                    continue

                distance_written = False
                for column, cell in vehicle_cells:
                    if self.stop_event.is_set():
                        break
                    self._wait_if_paused()
                    try:
                        assert column.rule is not None
                        requirements = "、".join(column.rule.vehicle_requirements) or "未指定"
                        self._event(
                            "log",
                            f"第 {row_idx} 行 {column.header}：开始报价，车长 {column.rule.car_length}，车型要求 {requirements}",
                        )
                        price = self._quote_vehicle_with_retries(client, column.rule)
                        write_success(cell, numeric_or_original(price))
                        self._event("log", f"第 {row_idx} 行 {column.header}：写入运费一口价 {price:g} 元")
                        if process_distance and not distance_written:
                            self._event("log", f"第 {row_idx} 行：正在读取总里程")
                            distance = client.read_distance()
                            write_success(distance_cell, numeric_or_original(distance))
                            distance_written = True
                            self._event(
                                "log",
                                f"第 {row_idx} 行：写入总里程 {distance:g} 公里到 {distance_col.letter} 列",
                            )
                    except Exception as exc:
                        message = self._failure_message(f"{column.header} 报价", exc)
                        mark_failure(cell, message)
                        self._event("log", f"第 {row_idx} 行 {column.header} 报价失败：{message}")

                if process_distance and not distance_written:
                    try:
                        first_rule = detection.vehicle_columns[0].rule
                        assert first_rule is not None
                        self._event("log", f"第 {row_idx} 行：价格列未成功写入路程，改用首个车型读取总里程")
                        self._quote_vehicle_with_retries(client, first_rule)
                        distance = client.read_distance()
                        write_success(distance_cell, numeric_or_original(distance))
                        self._event("log", f"第 {row_idx} 行：写入总里程 {distance:g} 公里到 {distance_col.letter} 列")
                    except Exception as exc:
                        message = self._failure_message("读取总里程", exc)
                        mark_failure(distance_cell, message)
                        self._event("log", f"第 {row_idx} 行：读取总里程失败：{message}")

                wb.save(output_path)
                self._event("log", f"第 {row_idx} 行：已保存到结果副本")
                self._event("progress", index, len(rows), f"{index}/{len(rows)}")

            self._write_address_mismatch_summary_sheet(wb, address_mismatches)
            if output_path:
                wb.save(output_path)
                self._log_address_mismatch_summary(address_mismatches)
                self._event("log", f"结果文件：{output_path}")
        except Exception:
            self._event("log", traceback.format_exc())
        finally:
            if client:
                client.stop()
            self._event("done", str(output_path) if output_path else "")

    def _wait_if_paused(self) -> None:
        while not self.pause_event.is_set():
            if self.stop_event.is_set():
                return
            time.sleep(0.2)

    def _quote_vehicle_with_retries(self, client: HuolalaClient, rule: VehicleRule, max_retries: int = 2) -> float:
        last_error: Exception | None = None
        for attempt in range(max_retries + 1):
            if self.stop_event.is_set():
                raise RuntimeError("任务已停止")
            try:
                return client.quote_vehicle(rule)
            except Exception as exc:
                last_error = exc
                if attempt >= max_retries:
                    break
                retry_index = attempt + 1
                self._event("log", f"{rule.name} 报价失败，准备第 {retry_index} 次重试：{exc}")
                time.sleep(0.8)
        raise RuntimeError(f"{rule.name} 报价失败，已重试 {max_retries} 次：{last_error}") from last_error

    def _compare_and_mark_addresses(
        self,
        ws,
        row_idx: int,
        supplier: str,
        origin_col,
        destination_col,
        expected_origin: str,
        expected_destination: str,
        actual_origin: str,
        actual_destination: str,
    ) -> list[dict[str, str | int]]:
        mismatches: list[dict[str, str | int]] = []
        checks = [
            ("发货地址", origin_col, expected_origin, actual_origin),
            ("到货地址", destination_col, expected_destination, actual_destination),
        ]
        for address_type, column, expected, actual in checks:
            expected_text = (expected or "").strip()
            actual_text = (actual or "").strip()
            cell = writable_cell(ws, row_idx, column.index)
            if expected_text == actual_text:
                clear_address_mismatch(cell)
                self._event("log", f"第 {row_idx} 行：{address_type}与 Excel 完全一致")
                continue

            actual_display = actual_text or "未读取到"
            expected_display = expected_text or "空"
            message = f"{address_type}不一致；Excel={expected_display}；网页参与计价={actual_display}"
            mark_address_mismatch(cell, message)
            mismatches.append(
                {
                    "row": row_idx,
                    "supplier": supplier or "",
                    "address_type": address_type,
                    "excel_address": expected_display,
                    "actual_address": actual_display,
                }
            )
            self._event(
                "log",
                f"第 {row_idx} 行：{address_type}与 Excel 不完全一致，需要人工核对；"
                f"Excel={expected_display}；网页参与计价={actual_display}",
            )
        return mismatches

    def _write_address_mismatch_summary_sheet(self, wb, mismatches: list[dict[str, str | int]]) -> None:
        sheet_name = "地址不一致汇总"
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        if not mismatches:
            return

        ws = wb.create_sheet(sheet_name)
        ws.append(["行号", "供应商", "地址类型", "Excel地址", "网页参与计价地址", "处理建议"])
        for item in mismatches:
            ws.append(
                [
                    item["row"],
                    item["supplier"],
                    item["address_type"],
                    item["excel_address"],
                    item["actual_address"],
                    "请人工核对并修正 Excel 地址后重跑",
                ]
            )
        for column, width in {
            "A": 10,
            "B": 28,
            "C": 14,
            "D": 46,
            "E": 46,
            "F": 30,
        }.items():
            ws.column_dimensions[column].width = width
        ws.freeze_panes = "A2"

    def _log_address_mismatch_summary(self, mismatches: list[dict[str, str | int]]) -> None:
        if not mismatches:
            self._event("log", "地址不一致汇总：0 项，网页实际参与计价地址与 Excel 地址完全一致。")
            return

        self._event(
            "log",
            f"地址不一致汇总：共 {len(mismatches)} 项，需要人工手动改正；"
            "已写入结果文件的“地址不一致汇总”工作表，并在原地址单元格加批注。",
        )
        for item in mismatches:
            self._event(
                "log",
                f"  第 {item['row']} 行 {item['supplier']} {item['address_type']}："
                f"Excel={item['excel_address']}；网页参与计价={item['actual_address']}",
            )

    def _failure_message(self, step: str, exc: Exception) -> str:
        raw = " ".join(str(exc).split())
        analysis = "原因待人工确认"
        if "地址识别" in raw or "intercepts pointer events" in raw or "选择准确地址" in raw:
            analysis = "地址识别候选没有匹配 Excel 中的供应商/发货地址/到货地址，或弹窗没有正常关闭"
        elif "无法选择车长" in raw:
            analysis = "网页没有成功切到目标车长，可能页面未加载完成、登录失效或车长文案变化"
        elif "车型要求未切换" in raw or "无法选择车型要求" in raw:
            analysis = "车型要求没有切换到规则指定状态，可能有错误勾选残留或网页限制至少保留一项"
        elif "运费一口价" in raw or "计价中" in raw:
            analysis = "网页没有生成可读取的运费一口价，可能计价仍在加载或当前路线暂不支持报价"
        elif "总里程" in raw:
            analysis = "网页没有生成可读取的总里程，可能地址未被准确识别"
        elif "Timeout" in raw:
            analysis = "网页响应超时，可能网络慢、登录过期或页面控件被弹窗挡住"
        detail = raw[:260]
        return f"{step}失败：{analysis}。原始提示：{detail}"

    def _ask_user(self, title: str, message: str) -> bool:
        result_queue: queue.Queue[bool] = queue.Queue(maxsize=1)

        def ask() -> None:
            result_queue.put(messagebox.askokcancel(title, message, parent=self.root))

        self.root.after(0, ask)
        return result_queue.get()

    def _event(self, kind: str, *payload) -> None:
        self.events.put((kind, *payload))

    def _drain_events(self) -> None:
        try:
            while True:
                event = self.events.get_nowait()
                kind = event[0]
                if kind == "log":
                    self._log(event[1])
                elif kind == "progress":
                    current, total, label = event[1], event[2], event[3]
                    value = 0 if total == 0 else current / total * 100
                    self.progress_var.set(value)
                    self.progress_label_var.set(label)
                elif kind == "result_path":
                    self._set_result_path(event[1])
                elif kind == "done":
                    self._set_running(False)
                    self.pause_button.configure(text="暂停")
                    output = event[1]
                    if output:
                        self._set_result_path(output)
                        self._refresh_result_buttons(running=False)
                        self.progress_label_var.set("已完成")
                    else:
                        self.progress_label_var.set("已结束")
        except queue.Empty:
            pass
        self.root.after(100, self._drain_events)

    def _log(self, message: str) -> None:
        timestamp = time.strftime("%H:%M:%S")
        self.log_text.configure(state="normal")
        self.log_text.insert("end", f"[{timestamp}] {message}\n")
        self.log_text.see("end")
        self.log_text.configure(state="disabled")
        with self.log_path.open("a", encoding="utf-8") as f:
            f.write(f"[{timestamp}] {message}\n")

    def _set_running(self, running: bool) -> None:
        can_start = bool(self.detection and self.detection.can_start)
        self.start_button.configure(state="disabled" if running or not can_start else "normal")
        self.pause_button.configure(state="normal" if running else "disabled")
        self.stop_button.configure(state="normal" if running else "disabled")
        self.strategy_combo.configure(state="disabled" if running else "readonly")
        self._refresh_result_buttons(running=running)

    def _set_result_path(self, path_text: str) -> None:
        self.last_output_path = Path(path_text) if path_text else None
        self.result_file_var.set(str(self.last_output_path) if self.last_output_path else "任务完成后会显示结果文件")
        running = bool(self.worker and self.worker.is_alive())
        self._refresh_result_buttons(running=running)

    def _refresh_result_buttons(self, running: bool) -> None:
        can_open = bool(self.last_output_path and self.last_output_path.exists() and not running)
        state = "normal" if can_open else "disabled"
        self.open_result_button.configure(state=state)
        self.open_result_folder_button.configure(state=state)

    def _open_result_file(self) -> None:
        if not self.last_output_path or not self.last_output_path.exists():
            messagebox.showwarning("结果文件不存在", "当前还没有可打开的结果文件。", parent=self.root)
            return
        self._open_path(self.last_output_path, "结果文件")

    def _open_result_folder(self) -> None:
        if not self.last_output_path or not self.last_output_path.exists():
            messagebox.showwarning("结果文件不存在", "当前还没有可打开的结果文件。", parent=self.root)
            return
        try:
            if sys.platform.startswith("win"):
                subprocess.Popen(["explorer", "/select,", str(self.last_output_path)])
            else:
                self._open_path(self.last_output_path.parent, "结果文件夹")
        except Exception as exc:
            messagebox.showerror("打开失败", f"无法打开结果文件夹：{exc}", parent=self.root)
            self._log(traceback.format_exc())

    def _open_path(self, path: Path, label: str) -> None:
        try:
            if sys.platform.startswith("win"):
                os.startfile(path)  # type: ignore[attr-defined]
            elif sys.platform == "darwin":
                subprocess.Popen(["open", str(path)])
            else:
                subprocess.Popen(["xdg-open", str(path)])
            self._log(f"已打开{label}：{path}")
        except Exception as exc:
            messagebox.showerror("打开失败", f"无法打开{label}：{exc}", parent=self.root)
            self._log(traceback.format_exc())

    def _bring_to_front(self) -> None:
        try:
            self.root.deiconify()
            self.root.state("normal")
            self.root.lift()
            self.root.focus_force()
            self.root.attributes("-topmost", True)
            self.root.after(1200, lambda: self.root.attributes("-topmost", False))
        except Exception:
            pass

    def _on_close(self) -> None:
        if self.worker and self.worker.is_alive():
            if not messagebox.askokcancel("退出", "任务仍在运行，确定要停止并退出吗？", parent=self.root):
                return
            self.stop_event.set()
            self.pause_event.set()
        self.root.destroy()


def main() -> None:
    root = tk.Tk()
    HuolalaQuoteApp(root)
    root.mainloop()
